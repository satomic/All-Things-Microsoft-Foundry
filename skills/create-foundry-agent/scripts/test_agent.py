"""
End-to-end test for a Microsoft Foundry (new API) Prompt Agent:
conversation -> Responses API call with agent_reference -> download files from
container_file_citation annotations -> optional .xlsx formula verification.

Usage:
    python test_agent.py --prompt-file test_prompt.txt          # agent that produces files
    python test_agent.py --prompt-file test_prompt.txt --no-files   # text-only agent

    Reads agent_info.json written by create_agent.py.

Dependencies:
    pip install "azure-ai-projects>=2.0.0" azure-identity openai
    pip install openpyxl   # only for the xlsx verification step
"""
import argparse
import json
from pathlib import Path

from azure.ai.projects import AIProjectClient
from azure.identity import DefaultAzureCredential

OUTPUT_DIR = Path(__file__).parent / "output"

# A good test prompt:
#  - contains all input data inline (no external sources needed)
#  - invokes the instructions' demo mode (no interactive confirmations)
#  - for file-producing agents, explicitly requests a downloadable artifact
DEFAULT_PROMPT = """\
Demo mode — produce the complete result end-to-end in one shot, no step-by-step confirmation.
<REPLACE WITH AN AGENT-SPECIFIC TEST REQUEST INCLUDING SAMPLE INPUT DATA>
"""


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prompt-file", default=None, help="file containing the test prompt")
    parser.add_argument("--no-files", action="store_true",
                        help="agent is not expected to produce downloadable files")
    args = parser.parse_args()
    prompt = Path(args.prompt_file).read_text(encoding="utf-8") if args.prompt_file else DEFAULT_PROMPT
    info = json.loads((Path(__file__).parent / "agent_info.json").read_text(encoding="utf-8"))

    with (
        DefaultAzureCredential() as credential,
        AIProjectClient(endpoint=info["endpoint"], credential=credential) as project_client,
        project_client.get_openai_client() as openai_client,
    ):
        conversation = openai_client.conversations.create()
        print(f"Conversation: {conversation.id}")

        print("Running agent via Responses API (may take a few minutes)...")
        response = openai_client.responses.create(
            conversation=conversation.id,
            input=prompt,
            extra_body={"agent_reference": {"name": info["agent_name"], "type": "agent_reference"}},
        )
        print(f"Response completed: id={response.id} status={response.status}")
        print("\n--- assistant ---")
        print(response.output_text[:3000])

        # Generated files live in container_file_citation annotations (NOT attachments)
        files = []  # (container_id, file_id, filename)
        for out in response.output:
            if out.type != "message":
                continue
            for content in out.content:
                if content.type != "output_text":
                    continue
                for ann in content.annotations or []:
                    if ann.type == "container_file_citation":
                        files.append((ann.container_id, ann.file_id, ann.filename))

        files = list(dict.fromkeys(files))
        print(f"\nGenerated files: {[f[2] for f in files]}")
        if args.no_files:
            print("RESULT: response completed (text-only agent; judge output against the "
                  "instruction contract)")
            return
        if not files:
            print("FAIL: no file generated")
            raise SystemExit(1)

        OUTPUT_DIR.mkdir(exist_ok=True)
        downloaded = []
        for container_id, file_id, filename in files:
            content = openai_client.containers.files.content.retrieve(
                file_id=file_id, container_id=container_id
            )
            target = OUTPUT_DIR / Path(filename).name
            target.write_bytes(content.read())
            print(f"Downloaded: {target} ({target.stat().st_size} bytes)")
            downloaded.append(target)

    for xlsx in (p for p in downloaded if p.suffix == ".xlsx"):
        verify_workbook(xlsx)


def verify_workbook(path: Path) -> None:
    """Spreadsheet-skill checks: derived values are formulas, statistics block, notes present."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    print(f"\n=== Verifying {path.name} ===")
    print(f"Sheets: {wb.sheetnames}")

    formulas, stat_funcs = [], set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
                    for fn in ("MEDIAN", "QUARTILE", "MAX", "MIN"):
                        if fn in cell.value:
                            stat_funcs.add(fn)

    checks = {
        "has formulas (not hardcoded)": len(formulas) >= 10,
        "has statistics block (MEDIAN/QUARTILE)": {"MEDIAN", "QUARTILE"} <= stat_funcs,
        "has notes/methodology": any("note" in s.lower() for s in wb.sheetnames),
    }

    print(f"Formula cells: {len(formulas)}; stat functions: {sorted(stat_funcs)}")
    for f in formulas[:10]:
        print(f"  {f}")

    all_pass = True
    for name, ok in checks.items():
        print(f"[{'PASS' if ok else 'FAIL'}] {name}")
        all_pass &= ok
    print("\nRESULT:", "ALL CHECKS PASSED" if all_pass else "SOME CHECKS FAILED")


if __name__ == "__main__":
    main()
